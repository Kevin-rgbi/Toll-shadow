"""Strict NYCCAS workbook quality and neighborhood coverage derivation."""

from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


class QualityDerivationError(ValueError):
    pass


WORKBOOKS = {
    "EC": {
        "sheet": "raw_data",
        "source_id": "nyccas_ec_year1_17_20260810",
        "label": "Elemental and black carbon",
        "fields": {"bc_conc": "ug/m3", "ec_abs_iso": "absorbance"},
        "start": "start_date",
        "end": "end_date",
        "headers": [
            "bc_conc", "ec_abs_iso", "start_date", "end_date", "season", "xcoordinate",
            "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
            "post_no", "session_id", "flag1", "flag2", "core",
        ],
    },
    "NOX": {
        "sheet": "raw_data",
        "source_id": "nyccas_nox_year1_17_20260810",
        "label": "Nitrogen oxides",
        "fields": {"blk_corr_no": "ppb", "blk_corr_no2": "ppb"},
        "start": "nox_start_datetime",
        "end": "nox_end_date_time",
        "headers": [
            "blk_corr_no", "blk_corr_no2", "nox_start_datetime", "nox_end_date_time", "season",
            "xcoordinate", "ycoordinate", "longitude", "latitude", "site_type", "Reference",
            "site_id", "post_no", "session_id", "flag1", "flag2", "core",
        ],
    },
    "PM": {
        "sheet": "PM_raw_data",
        "source_id": "nyccas_pm_year1_17_20260810",
        "label": "Fine particulate matter (PM2.5)",
        "fields": {"blk_corr_pm_ugm3": "ug/m3"},
        "start": "pm_start_date",
        "end": "pm_end_date",
        "headers": [
            "blk_corr_pm_ugm3", "pm_start_date", "pm_end_date", "season", "xcoordinate",
            "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
            "post_no", "session_id", "flag1", "flag2", "core",
        ],
    },
    "O3": {
        "sheet": "raw_data",
        "source_id": "nyccas_o3_year1_17_20260810",
        "label": "Ozone",
        "fields": {"blk_corr_o3_ppb": "ppb"},
        "start": "start_datetime",
        "end": "end_datetime",
        "headers": [
            "blk_corr_o3_ppb", "start_datetime", "end_datetime", "season", "xcoordinate",
            "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
            "post_no", "session_id", "flag1", "flag2", "core",
        ],
    },
}

LIMITATIONS = [
    "Raw NYCCAS samples are not directly representative of the sampled time or place.",
    "Temporal adjustment and modeling are required before describing pollution across New York City.",
    "These sampling methods are not comparable to regulatory monitoring.",
    "Counts describe source coverage and reported QA flags; they are not a statistical confidence score.",
]


def source_value(value: Any) -> Any | None:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value in {"", "NULL"}:
            return None
    return value


def as_datetime(value: Any, epoch: datetime, field: str) -> datetime:
    value = source_value(value)
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (int, float)) and math.isfinite(value):
        converted = from_excel(value, epoch)
        return converted if isinstance(converted, datetime) else datetime.combine(converted, time.min)
    raise QualityDerivationError(f"{field}: invalid Excel date {value!r}")


def finite_number(value: Any, field: str, *, minimum: float | None = None,
                  maximum: float | None = None) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise QualityDerivationError(f"{field}: invalid coordinate or numeric value {value!r}") from error
    if not math.isfinite(number) or (minimum is not None and number < minimum) or (maximum is not None and number > maximum):
        raise QualityDerivationError(f"{field}: coordinate or numeric value out of range {value!r}")
    return number


def _on_segment(point: tuple[float, float], left: tuple[float, float], right: tuple[float, float]) -> bool:
    px, py = point
    ax, ay = left
    bx, by = right
    cross = (px - ax) * (by - ay) - (py - ay) * (bx - ax)
    if abs(cross) > 1e-12:
        return False
    return min(ax, bx) - 1e-12 <= px <= max(ax, bx) + 1e-12 and min(ay, by) - 1e-12 <= py <= max(ay, by) + 1e-12


def _point_in_ring(point: tuple[float, float], ring: list[list[float]]) -> bool:
    inside = False
    for index, right in enumerate(ring):
        left = ring[index - 1]
        if _on_segment(point, tuple(left), tuple(right)):
            return True
        px, py = point
        ax, ay = left
        bx, by = right
        if (ay > py) != (by > py) and px < (bx - ax) * (py - ay) / (by - ay) + ax:
            inside = not inside
    return inside


def point_in_geometry(point: tuple[float, float], geometry: dict[str, Any]) -> bool:
    if geometry.get("type") != "MultiPolygon" or not isinstance(geometry.get("coordinates"), list):
        raise QualityDerivationError("Westchester Square geometry must be a MultiPolygon")
    for polygon in geometry["coordinates"]:
        if polygon and _point_in_ring(point, polygon[0]) and not any(_point_in_ring(point, hole) for hole in polygon[1:]):
            return True
    return False


def _point_segment_km(point: tuple[float, float], left: list[float], right: list[float]) -> float:
    lon, lat = point
    reference_latitude = math.radians((lat + left[1] + right[1]) / 3)
    x_scale = 111.32 * math.cos(reference_latitude)
    y_scale = 110.574
    px, py = lon * x_scale, lat * y_scale
    ax, ay = left[0] * x_scale, left[1] * y_scale
    bx, by = right[0] * x_scale, right[1] * y_scale
    dx, dy = bx - ax, by - ay
    ratio = 0.0 if dx == 0 and dy == 0 else max(0.0, min(1.0, ((px - ax) * dx + (py - ay) * dy) / (dx * dx + dy * dy)))
    return math.hypot(px - (ax + ratio * dx), py - (ay + ratio * dy))


def distance_to_geometry_km(point: tuple[float, float], geometry: dict[str, Any]) -> float:
    if point_in_geometry(point, geometry):
        return 0.0
    segments = (
        _point_segment_km(point, ring[index - 1], ring[index])
        for polygon in geometry["coordinates"]
        for ring in polygon
        for index in range(len(ring))
    )
    try:
        return min(segments)
    except ValueError as error:
        raise QualityDerivationError("Westchester Square geometry has no boundary segments") from error


def _read_pollutant(pollutant: str, path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    config = WORKBOOKS[pollutant]
    workbook = load_workbook(path, read_only=True, data_only=True)
    if config["sheet"] not in workbook.sheetnames:
        raise QualityDerivationError(f"{pollutant}: required sheet {config['sheet']!r} is missing")
    sheet = workbook[config["sheet"]]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as error:
        raise QualityDerivationError(f"{pollutant}: raw sheet is empty") from error
    if headers != config["headers"]:
        raise QualityDerivationError(f"{pollutant}: headers do not match the approved workbook contract")

    indexes = {name: index for index, name in enumerate(headers)}
    analytic = {name: {"unit": unit, "present": 0, "missing": 0} for name, unit in config["fields"].items()}
    seen: set[tuple[str, str, str]] = set()
    site_posts: dict[tuple[str, str], dict[str, Any]] = {}
    starts: list[datetime] = []
    qa = {"flag1": 0, "flag2": 0, "either_flag": 0, "neither_flag": 0}
    source_rows = 0

    for row_number, row_values in enumerate(rows, start=2):
        if all(source_value(value) is None for value in row_values):
            continue
        row = {name: row_values[index] if index < len(row_values) else None for name, index in indexes.items()}
        source_rows += 1
        site_id = str(source_value(row["site_id"]) or "").strip()
        post_no = str(source_value(row["post_no"]) or "").strip()
        if not site_id or not post_no:
            raise QualityDerivationError(f"{pollutant} row {row_number}: site_id and post_no are required")
        started = as_datetime(row[config["start"]], workbook.epoch, f"{pollutant} row {row_number} start")
        ended = as_datetime(row[config["end"]], workbook.epoch, f"{pollutant} row {row_number} end")
        if ended < started:
            raise QualityDerivationError(f"{pollutant} row {row_number}: end precedes start")
        key = (site_id, post_no, started.isoformat())
        if key in seen:
            raise QualityDerivationError(f"{pollutant} row {row_number}: duplicate site/post/start key {key}")
        seen.add(key)
        starts.append(started)

        longitude = finite_number(row["longitude"], f"{pollutant} row {row_number} longitude", minimum=-180, maximum=180)
        latitude = finite_number(row["latitude"], f"{pollutant} row {row_number} latitude", minimum=-90, maximum=90)
        for field in config["fields"]:
            value = source_value(row[field])
            if value is None:
                analytic[field]["missing"] += 1
            else:
                finite_number(value, f"{pollutant} row {row_number} {field}")
                analytic[field]["present"] += 1

        flag1 = source_value(row["flag1"])
        flag2 = source_value(row["flag2"])
        qa["flag1"] += int(flag1 is not None)
        qa["flag2"] += int(flag2 is not None)
        qa["either_flag"] += int(flag1 is not None or flag2 is not None)
        qa["neither_flag"] += int(flag1 is None and flag2 is None)

        site_key = (site_id, post_no)
        item = site_posts.get(site_key)
        if item is None:
            item = {
                "site_id": site_id,
                "post_no": post_no,
                "latitude": latitude,
                "longitude": longitude,
                "reference_values": set(),
                "core_values": set(),
                "source_rows": 0,
                "qa_rows": 0,
                "starts": [],
            }
            site_posts[site_key] = item
        elif abs(item["latitude"] - latitude) > 1e-6 or abs(item["longitude"] - longitude) > 1e-6:
            raise QualityDerivationError(f"{pollutant} row {row_number}: site/post coordinate changed")
        item["source_rows"] += 1
        item["qa_rows"] += int(flag1 is not None or flag2 is not None)
        item["starts"].append(started)
        item["reference_values"].add(str(source_value(row["Reference"])))
        item["core_values"].add(str(source_value(row["core"])))

    if source_rows == 0:
        raise QualityDerivationError(f"{pollutant}: no source rows")

    published_sites = []
    locations = []
    for item in sorted(site_posts.values(), key=lambda value: (value["site_id"], value["post_no"])):
        coverage = {"start": min(item["starts"]).date().isoformat(), "end": max(item["starts"]).date().isoformat()}
        published_sites.append({
            "site_id": item["site_id"],
            "post_no": item["post_no"],
            "latitude": item["latitude"],
            "longitude": item["longitude"],
            "reference_values": sorted(item["reference_values"]),
            "core_values": sorted(item["core_values"]),
            "source_rows": item["source_rows"],
            "qa_rows": item["qa_rows"],
            "coverage": coverage,
        })
        locations.append({
            "site_id": item["site_id"], "post_no": item["post_no"], "latitude": item["latitude"],
            "longitude": item["longitude"], "pollutant": pollutant,
        })

    return ({
        "source_id": config["source_id"],
        "label": config["label"],
        "source_rows": source_rows,
        "distinct_site_ids": len({item["site_id"] for item in published_sites}),
        "distinct_site_posts": len(published_sites),
        "coverage": {"start": min(starts).date().isoformat(), "end": max(starts).date().isoformat()},
        "analytic_fields": analytic,
        "qa": qa,
        "site_posts": published_sites,
    }, locations)


def _read_nta(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("type") != "FeatureCollection" or len(payload.get("features", [])) != 1:
        raise QualityDerivationError("Westchester Square source must contain exactly one feature")
    feature = payload["features"][0]
    properties = feature.get("properties", {})
    if properties.get("nta2020") != "BX1001" or properties.get("ntaname") != "Westchester Square":
        raise QualityDerivationError("Westchester Square source is not NTA BX1001")
    point_in_geometry((-73.0, 40.0), feature.get("geometry", {}))
    return feature


def _current_monitors(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        required = {"date_nyc", "site_id", "site_name", "borough", "latitude", "longitude"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise QualityDerivationError("current-air CSV is missing required headers")
        latest: dict[str, dict[str, Any]] = {}
        for row_number, row in enumerate(reader, start=2):
            site_id = (row.get("site_id") or "").strip()
            day = (row.get("date_nyc") or "").strip()
            if not site_id or not day:
                raise QualityDerivationError(f"current-air row {row_number}: site_id and date_nyc are required")
            item = {
                "site_id": site_id,
                "site_name": (row.get("site_name") or site_id).strip(),
                "borough": (row.get("borough") or "Unknown").strip(),
                "latitude": finite_number(row.get("latitude"), f"current-air row {row_number} latitude", minimum=-90, maximum=90),
                "longitude": finite_number(row.get("longitude"), f"current-air row {row_number} longitude", minimum=-180, maximum=180),
                "date": day,
            }
            if site_id not in latest or day > latest[site_id]["date"]:
                latest[site_id] = item
    if not latest:
        raise QualityDerivationError("current-air CSV contains no monitor rows")
    return list(latest.values())


def derive_quality(workbooks: dict[str, Path], nta_path: Path, current_air_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    if set(workbooks) != set(WORKBOOKS):
        raise QualityDerivationError(f"workbooks must be exactly {sorted(WORKBOOKS)}")
    pollutants = {}
    all_locations: dict[tuple[str, str, float, float], dict[str, Any]] = {}
    for pollutant in WORKBOOKS:
        summary, locations = _read_pollutant(pollutant, Path(workbooks[pollutant]))
        pollutants[pollutant] = summary
        for location in locations:
            key = (location["site_id"], location["post_no"], location["latitude"], location["longitude"])
            existing = all_locations.setdefault(key, {**location, "pollutants": []})
            existing["pollutants"].append(pollutant)
            existing.pop("pollutant", None)

    coverage_starts = [summary["coverage"]["start"] for summary in pollutants.values()]
    coverage_ends = [summary["coverage"]["end"] for summary in pollutants.values()]
    quality = {
        "schema_version": "1.0.0",
        "title": "NYCCAS year 1 through 17 data quality and coverage",
        "source_release_date": "2026-08-10",
        "coverage": {"start": min(coverage_starts), "end": max(coverage_ends)},
        "pollutants": pollutants,
        "limitations": LIMITATIONS,
    }

    boundary = _read_nta(Path(nta_path))
    geometry = boundary["geometry"]
    historical = list(all_locations.values())
    current = _current_monitors(Path(current_air_path))
    for item in historical + current:
        point = (item["longitude"], item["latitude"])
        item["inside"] = point_in_geometry(point, geometry)
        item["distance_to_boundary_km"] = round(distance_to_geometry_km(point, geometry), 3)

    nearest_historical = min(historical, key=lambda item: (item["distance_to_boundary_km"], item["site_id"], item.get("post_no", "")))
    nearest_current = min(current, key=lambda item: (item["distance_to_boundary_km"], item["site_id"]))
    boundary_feature = {
        "type": "Feature",
        "geometry": geometry,
        "properties": {**boundary["properties"], "kind": "boundary"},
    }

    def point_feature(kind: str, item: dict[str, Any]) -> dict[str, Any]:
        properties = {key: value for key, value in item.items() if key not in {"latitude", "longitude", "date"}}
        return {
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [item["longitude"], item["latitude"]]},
            "properties": {"kind": kind, **properties},
        }

    neighborhood = {
        "type": "FeatureCollection",
        "schema_version": "1.0.0",
        "geometry_crs": "EPSG:4326",
        "name": "Westchester Square source coverage",
        "coverage_summary": {
            "historical_sites_inside": sum(item["inside"] for item in historical),
            "current_monitors_inside": sum(item["inside"] for item in current),
        },
        "features": [
            boundary_feature,
            point_feature("nearest_historical", nearest_historical),
            point_feature("nearest_current", nearest_current),
        ],
        "limitations": [
            "The official NTA is a statistical geography and may not definitively represent the neighborhood.",
            "Nearest outside sites are not Westchester Square measurements and are not used to estimate a neighborhood value.",
        ],
    }
    return quality, neighborhood
