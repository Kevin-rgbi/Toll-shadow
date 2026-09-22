from __future__ import annotations

import csv
import json
import math
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from pipeline.src.nyccas_quality import (
    QualityDerivationError,
    derive_quality,
    distance_to_geometry_km,
    point_in_geometry,
)


HEADERS = {
    "EC": [
        "bc_conc", "ec_abs_iso", "start_date", "end_date", "season", "xcoordinate",
        "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
        "post_no", "session_id", "flag1", "flag2", "core",
    ],
    "NOX": [
        "blk_corr_no", "blk_corr_no2", "nox_start_datetime", "nox_end_date_time", "season",
        "xcoordinate", "ycoordinate", "longitude", "latitude", "site_type", "Reference",
        "site_id", "post_no", "session_id", "flag1", "flag2", "core",
    ],
    "PM": [
        "blk_corr_pm_ugm3", "pm_start_date", "pm_end_date", "season", "xcoordinate",
        "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
        "post_no", "session_id", "flag1", "flag2", "core",
    ],
    "O3": [
        "blk_corr_o3_ppb", "start_datetime", "end_datetime", "season", "xcoordinate",
        "ycoordinate", "longitude", "latitude", "site_type", "Reference", "site_id",
        "post_no", "session_id", "flag1", "flag2", "core",
    ],
}


def row_for(pollutant: str, *, site_id: str = "outside", longitude: float = 2.0,
            latitude: float = 0.5, flag1: str | None = None, flag2: str | None = None,
            first_value: float | str | None = 5.0) -> list[object]:
    values = {
        "bc_conc": first_value,
        "ec_abs_iso": 1.25,
        "blk_corr_no": first_value,
        "blk_corr_no2": 10.0,
        "blk_corr_pm_ugm3": first_value,
        "blk_corr_o3_ppb": first_value,
        "start_date": datetime(2025, 1, 1),
        "end_date": datetime(2025, 1, 14),
        "nox_start_datetime": datetime(2025, 1, 1, 8),
        "nox_end_date_time": datetime(2025, 1, 14, 8),
        "pm_start_date": datetime(2025, 1, 1),
        "pm_end_date": datetime(2025, 1, 14),
        "start_datetime": datetime(2025, 6, 1, 8),
        "end_datetime": datetime(2025, 6, 14, 8),
        "season": "W17",
        "xcoordinate": 1,
        "ycoordinate": 1,
        "longitude": longitude,
        "latitude": latitude,
        "site_type": "Street",
        "Reference": 0,
        "site_id": site_id,
        "post_no": 1,
        "session_id": 101,
        "flag1": flag1,
        "flag2": flag2,
        "core": 2,
    }
    return [values[name] for name in HEADERS[pollutant]]


def write_workbook(path: Path, pollutant: str, rows: list[list[object]], *, raw_sheet: str | None = None,
                   headers: list[str] | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = raw_sheet or ("PM_raw_data" if pollutant == "PM" else "raw_data")
    sheet.append(headers or HEADERS[pollutant])
    for row in rows:
        sheet.append(row)
    workbook.create_sheet("data dictionary")
    workbook.create_sheet("metadata")
    workbook.save(path)


def write_nta(path: Path) -> None:
    payload = {
        "type": "FeatureCollection",
        "features": [{
            "type": "Feature",
            "properties": {"nta2020": "BX1001", "ntaname": "Westchester Square", "boroname": "Bronx"},
            "geometry": {
                "type": "MultiPolygon",
                "coordinates": [[
                    [[0, 0], [1, 0], [1, 1], [0, 1], [0, 0]],
                    [[0.4, 0.4], [0.6, 0.4], [0.6, 0.6], [0.4, 0.6], [0.4, 0.4]],
                ]],
            },
        }],
    }
    path.write_text(json.dumps(payload), encoding="utf-8")


def write_current_air(path: Path, rows: list[dict[str, str]] | None = None) -> None:
    fields = ["date_nyc", "site_id", "site_name", "borough", "latitude", "longitude"]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows or [{
            "date_nyc": "2026-09-20", "site_id": "current-outside", "site_name": "Outside monitor",
            "borough": "Bronx", "latitude": "0.5", "longitude": "3.0",
        }]:
            writer.writerow(row)


class NyccasQualityTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workbooks: dict[str, Path] = {}
        for pollutant in HEADERS:
            path = self.root / f"{pollutant}.xlsx"
            write_workbook(path, pollutant, [row_for(pollutant)])
            self.workbooks[pollutant] = path
        self.nta = self.root / "nta.geojson"
        self.current_air = self.root / "current.csv"
        write_nta(self.nta)
        write_current_air(self.current_air)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_counts_nulls_and_flags_without_excluding_rows(self) -> None:
        write_workbook(self.workbooks["PM"], "PM", [
            row_for("PM", site_id="a", first_value=5.0),
            row_for("PM", site_id="b", first_value="NULL", flag1="HF"),
            row_for("PM", site_id="c", first_value=None, flag2="LF"),
        ])

        quality, _ = derive_quality(self.workbooks, self.nta, self.current_air)
        pm = quality["pollutants"]["PM"]

        self.assertEqual(pm["source_rows"], 3)
        self.assertEqual(pm["analytic_fields"]["blk_corr_pm_ugm3"]["present"], 1)
        self.assertEqual(pm["analytic_fields"]["blk_corr_pm_ugm3"]["missing"], 2)
        self.assertEqual(pm["qa"], {"flag1": 1, "flag2": 1, "either_flag": 2, "neither_flag": 1})
        self.assertEqual(len(pm["site_posts"]), 3)
        self.assertNotIn("blk_corr_pm_ugm3", pm["site_posts"][0])

    def test_keeps_both_ec_analytic_fields_separate(self) -> None:
        write_workbook(self.workbooks["EC"], "EC", [
            row_for("EC", first_value="NULL"),
        ])
        quality, _ = derive_quality(self.workbooks, self.nta, self.current_air)
        fields = quality["pollutants"]["EC"]["analytic_fields"]
        self.assertEqual(fields["bc_conc"]["missing"], 1)
        self.assertEqual(fields["ec_abs_iso"]["present"], 1)

    def test_rejects_missing_headers_and_duplicate_site_post_start(self) -> None:
        write_workbook(self.workbooks["PM"], "PM", [row_for("PM")], headers=HEADERS["PM"][:-1])
        with self.assertRaisesRegex(QualityDerivationError, "headers"):
            derive_quality(self.workbooks, self.nta, self.current_air)

        duplicate = row_for("PM")
        write_workbook(self.workbooks["PM"], "PM", [duplicate, duplicate])
        with self.assertRaisesRegex(QualityDerivationError, "duplicate"):
            derive_quality(self.workbooks, self.nta, self.current_air)

    def test_rejects_non_finite_or_out_of_range_coordinates(self) -> None:
        write_workbook(self.workbooks["O3"], "O3", [row_for("O3", longitude=math.inf)])
        with self.assertRaisesRegex(QualityDerivationError, "coordinate"):
            derive_quality(self.workbooks, self.nta, self.current_air)

    def test_geometry_holes_and_boundary_distance_are_honest(self) -> None:
        geometry = json.loads(self.nta.read_text())["features"][0]["geometry"]
        self.assertTrue(point_in_geometry((0.2, 0.2), geometry))
        self.assertFalse(point_in_geometry((0.5, 0.5), geometry))
        self.assertFalse(point_in_geometry((2.0, 0.5), geometry))
        self.assertAlmostEqual(distance_to_geometry_km((2.0, 0.5), geometry), 111.3, delta=1.0)

    def test_emits_boundary_and_nearest_outside_points(self) -> None:
        quality, neighborhood = derive_quality(self.workbooks, self.nta, self.current_air)
        self.assertEqual(quality["schema_version"], "1.0.0")
        self.assertEqual(neighborhood["coverage_summary"], {
            "historical_sites_inside": 0,
            "current_monitors_inside": 0,
        })
        self.assertEqual([feature["properties"]["kind"] for feature in neighborhood["features"]], [
            "boundary", "nearest_historical", "nearest_current",
        ])
        for feature in neighborhood["features"][1:]:
            self.assertFalse(feature["properties"]["inside"])
            self.assertGreater(feature["properties"]["distance_to_boundary_km"], 0)


if __name__ == "__main__":
    unittest.main()
